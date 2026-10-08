from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula


def is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def find_header_row(worksheet, search_limit: int = 50) -> int | None:
    for row_index in range(1, min(worksheet.max_row, search_limit) + 1):
        if any(
            isinstance(worksheet.cell(row_index, column_index).value, str)
            and worksheet.cell(row_index, column_index).value.strip() == "#"
            for column_index in range(1, worksheet.max_column + 1)
        ):
            return row_index
    return None


def _format_value(value: object, formula_value: object) -> str:
    if is_blank(value):
        if isinstance(formula_value, ArrayFormula):
            return formula_value.text or ""
        if isinstance(formula_value, str) and formula_value.startswith("="):
            return formula_value
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value).strip() if isinstance(value, str) else str(value)


def _unique_headers(worksheet, header_row: int) -> list[tuple[int, str]]:
    counts: dict[str, int] = {}
    headers = []
    for column_index in range(1, worksheet.max_column + 1):
        value = worksheet.cell(header_row, column_index).value
        if is_blank(value):
            continue
        header = str(value).strip().replace("\n", " ")
        counts[header] = counts.get(header, 0) + 1
        if counts[header] > 1:
            header = f"{header} ({counts[header]})"
        headers.append((column_index, header))
    return headers


def create_markdown(
    workbook_path: Path,
    sheet_name: str | None = None,
    start_row: int | None = None,
) -> tuple[Path, int, str]:
    workbook_path = workbook_path.expanduser().resolve()
    formulas = load_workbook(workbook_path, data_only=False)
    values = load_workbook(workbook_path, data_only=True)
    try:
        if sheet_name:
            if sheet_name not in formulas.sheetnames:
                raise ValueError(f"Worksheet not found: {sheet_name}")
            formula_sheet = formulas[sheet_name]
        else:
            formula_sheet = next(
                (sheet for sheet in formulas.worksheets if find_header_row(sheet) is not None),
                None,
            )
            if formula_sheet is None:
                raise ValueError("No worksheet with a '#' header row was found.")
        value_sheet = values[formula_sheet.title]
        header_row = find_header_row(formula_sheet)
        if header_row is None:
            raise ValueError(f"No '#' header row found in worksheet: {formula_sheet.title}")

        data_start_row = start_row if start_row is not None else header_row + 1
        if data_start_row <= header_row:
            raise ValueError(f"Conversion start row must be greater than the header row ({header_row}).")

        headers = _unique_headers(formula_sheet, header_row)
        blocks = []
        for row_index in range(data_start_row, formula_sheet.max_row + 1):
            fields = []
            for column_index, header in headers:
                formula_value = formula_sheet.cell(row_index, column_index).value
                value = value_sheet.cell(row_index, column_index).value
                text = _format_value(value, formula_value)
                if text:
                    fields.append(f"### {header}\n\n{text}")
            if fields:
                blocks.append("\n\n".join(fields))

        markdown_path = workbook_path.with_name(f"{formula_sheet.title}.md")
        content = f"# {formula_sheet.title}\n\n## Summary\n\n## Steps"
        if blocks:
            content += "\n\n" + "\n\n---\n\n".join(blocks)
        content += "\n"
        return markdown_path, len(blocks), content
    finally:
        formulas.close()
        values.close()
