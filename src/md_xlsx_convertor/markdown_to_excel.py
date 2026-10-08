from __future__ import annotations

import importlib
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from .excel_to_markdown import find_header_row, is_blank


def title_from_markdown(markdown_path: Path) -> str:
    for line in markdown_path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"^#\s+(.+?)\s*$", line)
        if match:
            return match.group(1)
    raise ValueError(f"Markdown has no H1 title: {markdown_path}")


def normalize_header(value: object) -> str:
    header = str(value).strip()
    header = re.sub(r"\s*[（(][^（）()]*[）)]$", "", header)
    return re.sub(r"\s+", "", header)


def numbered_headers(worksheet, header_row: int) -> dict[int, str]:
    counts: dict[str, int] = {}
    headers = {}
    for column_index in range(1, worksheet.max_column + 1):
        value = worksheet.cell(header_row, column_index).value
        if is_blank(value):
            continue
        base = normalize_header(value)
        counts[base] = counts.get(base, 0) + 1
        headers[column_index] = f"{base}#{counts[base]}"
    return headers


def resolve_source_workbook(
    markdown_path: Path,
    title: str,
    explicit_path: Path | None = None,
    preferred_sheet: str | None = None,
) -> tuple[Path, str]:
    candidates = [explicit_path.expanduser().resolve()] if explicit_path else sorted(markdown_path.parent.glob("*.xlsx"))
    matches = []
    for workbook_path in candidates:
        if not workbook_path.is_file():
            continue
        try:
            workbook = load_workbook(workbook_path, read_only=True, data_only=False)
            try:
                requested_sheet = preferred_sheet or title
                sheet_name = requested_sheet if requested_sheet in workbook.sheetnames else None
                if sheet_name is None:
                    sheet_name = next(
                        (name for name in workbook.sheetnames if find_header_row(workbook[name]) is not None),
                        None,
                    )
                if sheet_name and find_header_row(workbook[sheet_name]) is not None:
                    matches.append((workbook_path, sheet_name))
            finally:
                workbook.close()
        except (OSError, ValueError):
            continue

    if len(matches) != 1:
        if not matches:
            raise ValueError("No matching .xlsx with a '#' table header was found.")
        raise ValueError("Multiple matching .xlsx files were found. Specify the target workbook.")
    return matches[0]


def _find_generated_header_row(worksheet, source_headers: dict[int, str]) -> int:
    expected = set(source_headers.values())
    candidates: list[tuple[int, int]] = []
    for row_index in range(1, min(worksheet.max_row, 60) + 1):
        keys: set[str] = set()
        counts: dict[str, int] = {}
        for column_index in range(1, worksheet.max_column + 1):
            value = worksheet.cell(row_index, column_index).value
            if is_blank(value):
                continue
            base = normalize_header(value)
            counts[base] = counts.get(base, 0) + 1
            keys.add(f"{base}#{counts[base]}")
        matching = len(expected.intersection(keys))
        if matching:
            candidates.append((matching, -row_index))
    if not candidates:
        raise ValueError("Could not locate the generated table header row.")
    matching, negative_row = max(candidates)
    if matching < max(1, (len(expected) + 1) // 2):
        raise ValueError("Generated table headers do not sufficiently match the target worksheet.")
    return -negative_row


def _generated_rows(worksheet, header_row: int) -> list[dict[str, object]]:
    headers = numbered_headers(worksheet, header_row)
    rows = []
    for row_index in range(header_row + 1, worksheet.max_row + 1):
        row = {key: worksheet.cell(row_index, column).value for column, key in headers.items()}
        if any(not is_blank(value) for value in row.values()):
            rows.append(row)
    return rows


def _cell_value(formula_cell, cached_cell):
    if cached_cell.value is None and isinstance(formula_cell.value, ArrayFormula):
        return formula_cell.value.text
    return cached_cell.value


def update_workbook(
    source_path: Path,
    source_sheet_name: str,
    generated_path: Path,
    generated_sheet_name: str,
    start_row: int | None = None,
) -> int:
    workbook = load_workbook(source_path, data_only=False)
    cached_workbook = load_workbook(source_path, data_only=True, read_only=True)
    generated_workbook = load_workbook(generated_path, data_only=False, read_only=True)
    temporary_path: Path | None = None
    try:
        source_sheet = workbook[source_sheet_name]
        cached_sheet = cached_workbook[source_sheet_name]
        generated_sheet = generated_workbook[generated_sheet_name]
        source_header_row = find_header_row(source_sheet)
        if source_header_row is None:
            raise ValueError(f"No '#' table header in target worksheet: {source_sheet_name}")
        source_headers = numbered_headers(source_sheet, source_header_row)
        generated_header_row = _find_generated_header_row(generated_sheet, source_headers)
        generated_headers = numbered_headers(generated_sheet, generated_header_row)
        generated_rows = _generated_rows(generated_sheet, generated_header_row)
        target_columns = {key: column for column, key in source_headers.items()}
        unmatched = [key for key in generated_headers.values() if key not in target_columns]
        if unmatched:
            raise ValueError("Generated columns do not match target columns: " + ", ".join(unmatched))

        target_start = start_row if start_row is not None else source_header_row + 1
        if target_start <= source_header_row:
            raise ValueError(f"Copy start row must be greater than the header row ({source_header_row}).")

        number_column = next((column for column, key in source_headers.items() if key.startswith("##")), None)
        copied_columns = [column for column in target_columns.values() if column != number_column]
        last_existing_row = target_start - 1
        for row_index in range(target_start, source_sheet.max_row + 1):
            if any(
                not is_blank(_cell_value(source_sheet.cell(row_index, column), cached_sheet.cell(row_index, column)))
                for column in target_columns.values()
            ):
                last_existing_row = row_index
        clear_end = max(last_existing_row, target_start + len(generated_rows) - 1)
        for row_index in range(target_start, clear_end + 1):
            for column in copied_columns:
                source_sheet.cell(row_index, column).value = None

        for offset, row in enumerate(generated_rows):
            target_row = target_start + offset
            for header_key, value in row.items():
                target_column = target_columns[header_key]
                if target_column != number_column:
                    source_sheet.cell(target_row, target_column).value = value

        with tempfile.NamedTemporaryFile(
            prefix=source_path.stem + "_", suffix=".xlsx", dir=source_path.parent, delete=False
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        row_count = len(generated_rows)
    except BaseException:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
        cached_workbook.close()
        generated_workbook.close()

    try:
        temporary_path.replace(source_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return row_count


def ensure_mael(progress=None):
    try:
        return importlib.import_module("mael.excel_builder").convert
    except ImportError:
        if getattr(sys, "frozen", False):
            raise RuntimeError("mael is missing from this executable. Rebuild the packaged application.")
        if progress:
            progress("mael is missing; installing the pinned package...")
        result = subprocess.run(
            [sys.executable, "-m", "pip", "install", "mael==0.0.3.35"],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode:
            raise RuntimeError(result.stderr.strip() or "Could not install mael.")
        importlib.invalidate_caches()
        return importlib.import_module("mael.excel_builder").convert


def process_markdown(
    markdown_path: Path,
    source_workbook: Path | None = None,
    sheet_name: str | None = None,
    start_row: int | None = None,
    progress=None,
) -> tuple[Path, int]:
    markdown_path = markdown_path.expanduser().resolve()
    if not markdown_path.is_file() or markdown_path.suffix.lower() != ".md":
        raise ValueError(f"Markdown file not found: {markdown_path}")
    title = title_from_markdown(markdown_path)
    source_path, auto_sheet = resolve_source_workbook(markdown_path, title, source_workbook, sheet_name)
    source_sheet_name = sheet_name or auto_sheet
    directory = markdown_path.parent
    output_directory = directory / "output"
    expected_output = output_directory / f"{directory.name}.xlsx"
    if output_directory.exists():
        contents = list(output_directory.iterdir())
        unexpected = [path for path in contents if path.resolve() != expected_output.resolve()]
        if unexpected:
            raise ValueError("Refusing to run because output contains unrelated files or folders.")
        if contents and not expected_output.is_file():
            raise ValueError("Refusing to overwrite a non-file mael output path.")

    if progress:
        progress("Building workbook with mael...")
    ensure_mael(progress)(str(directory))
    if not expected_output.is_file():
        raise FileNotFoundError(f"mael output not found: {expected_output}")
    generated = load_workbook(expected_output, read_only=True, data_only=False)
    try:
        if title not in generated.sheetnames:
            raise ValueError(f"mael output has no worksheet named '{title}'")
    finally:
        generated.close()

    if progress:
        progress("Updating the target workbook...")
    row_count = update_workbook(source_path, source_sheet_name, expected_output, title, start_row)
    items = list(output_directory.iterdir())
    if items != [expected_output]:
        raise ValueError("Workbook updated, but output contains unrelated items and was retained.")
    shutil.rmtree(output_directory)
    return source_path, row_count
