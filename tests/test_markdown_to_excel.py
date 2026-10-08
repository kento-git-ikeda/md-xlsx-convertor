from __future__ import annotations

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from md_xlsx_convertor.markdown_to_excel import process_markdown


class MarkdownToExcelTests(unittest.TestCase):
    def test_builds_and_updates_existing_workbook_from_requested_row(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            markdown = root / "Cases.md"
            markdown.write_text(
                "# Cases\n\n## Summary\n\n## Steps\n\n"
                "### #\n\n1\n\n### Name\n\nnew value\n",
                encoding="utf-8",
            )
            workbook_path = root / "spec.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Cases"
            sheet.append(["Cover"])
            sheet.append(["#", "Name"])
            sheet.append([10, "keep above"])
            sheet.append([11, "replace me"])
            workbook.create_sheet("Summary")
            workbook.save(workbook_path)
            workbook.close()

            updated_path, row_count = process_markdown(markdown, workbook_path, start_row=4)

            result = load_workbook(updated_path, data_only=False)
            try:
                self.assertEqual(row_count, 1)
                self.assertEqual(result.sheetnames, ["Cases", "Summary"])
                self.assertEqual(result["Cases"]["B3"].value, "keep above")
                self.assertEqual(result["Cases"]["B4"].value, "new value")
                self.assertFalse((root / "output").exists())
            finally:
                result.close()

    def test_refuses_to_remove_unrelated_output_files(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            markdown = root / "Cases.md"
            markdown.write_text("# Cases\n\n## Summary\n\n## Steps\n", encoding="utf-8")
            output = root / "output"
            output.mkdir()
            (output / "keep.txt").write_text("keep", encoding="utf-8")
            with self.assertRaises(ValueError):
                process_markdown(markdown)
            self.assertEqual((output / "keep.txt").read_text(encoding="utf-8"), "keep")


if __name__ == "__main__":
    unittest.main()
